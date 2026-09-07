"""
v0.1.94：矿山设备资料生成Word文档导出

将生成的工程资料导出为标准Word文档，支持8种现场记录类型的Word模板，
自动填充数据到Word文档，支持下载。
"""

import os
import datetime
from typing import Optional

try:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


# ============================================================
# Word文档导出
# ============================================================

def export_record_to_word(
    record_type: str,
    record_data: dict,
    device_type: str = "",
    device_tag: str = "",
    workshop: str = "",
    output_path: str = "",
) -> dict:
    """
    将现场记录导出为Word文档。
    
    Args:
        record_type: 记录类型
        record_data: 记录数据
        device_type: 设备类型
        device_tag: 设备位号
        workshop: 车间
        output_path: 输出路径
    
    Returns:
        导出结果
    """
    if not HAS_DOCX:
        return {"ok": False, "error": "python-docx 未安装，无法导出Word文档"}
    
    # 获取记录类型名称
    type_names = {
        "construction_log": "施工日志",
        "unboxing_record": "设备开箱检验记录",
        "concealment_record": "隐蔽工程验收记录",
        "installation_record": "设备安装记录",
        "trial_run_record": "设备试运转记录",
        "safety_check": "安全检查记录",
        "lifting_record": "吊装作业记录",
        "welding_record": "焊接记录",
    }
    
    doc_name = type_names.get(record_type, "工程资料")
    
    # 创建文档
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # 标题
    title = doc.add_heading(doc_name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = '黑体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0, 0, 0)
    
    # 设备信息
    if device_type or device_tag or workshop:
        info_para = doc.add_paragraph()
        info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        info_text = ""
        if device_type:
            info_text += f"设备名称：{device_type}    "
        if device_tag:
            info_text += f"设备位号：{device_tag}    "
        if workshop:
            info_text += f"施工区域：{workshop}"
        run = info_para.add_run(info_text)
        run.font.size = Pt(11)
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    doc.add_paragraph()  # 空行
    
    # 获取字段映射
    field_mappings = _get_field_mapping(record_type)
    
    # 创建表格
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 设置表格列宽
    for row in table.rows:
        row.cells[0].width = Cm(5)
        row.cells[1].width = Cm(11)
    
    # 表头
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '项目'
    hdr_cells[1].text = '内容'
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.name = '黑体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(11)
    
    # 填充数据
    for field_key, field_label in field_mappings:
        value = record_data.get(field_key, "")
        if not value or value == "待填写":
            value = "________________"
        
        row_cells = table.add_row().cells
        row_cells[0].text = field_label
        row_cells[1].text = str(value)
        
        # 设置字体
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = '宋体'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                    run.font.size = Pt(11)
    
    # 设备专用要点
    equipment_points = _get_equipment_points(record_type, device_type)
    if equipment_points:
        doc.add_paragraph()
        points_heading = doc.add_heading('设备专用要点', level=2)
        for run in points_heading.runs:
            run.font.name = '黑体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            run.font.size = Pt(14)
        
        for point in equipment_points:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(point)
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(11)
    
    # 签字栏
    doc.add_paragraph()
    doc.add_paragraph()
    
    sign_table = doc.add_table(rows=2, cols=4)
    sign_table.style = 'Table Grid'
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    sign_labels = ['编制', '审核', '批准', '日期']
    for i, label in enumerate(sign_labels):
        sign_table.rows[0].cells[i].text = label
        sign_table.rows[1].cells[i].text = '________________'
        for cell in [sign_table.rows[0].cells[i], sign_table.rows[1].cells[i]]:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = '宋体'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                    run.font.size = Pt(11)
            if cell == sign_table.rows[0].cells[i]:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
    
    # 保存文档
    if not output_path:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = doc_name.replace("/", "_").replace(" ", "")
        output_path = f"/tmp/{safe_name}_{timestamp}.docx"
    
    # 确保目录存在
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    
    doc.save(output_path)
    
    file_size = os.path.getsize(output_path)
    
    return {
        "ok": True,
        "record_type": record_type,
        "doc_name": doc_name,
        "output_path": output_path,
        "file_size": file_size,
        "file_size_kb": round(file_size / 1024, 1),
        "field_count": len(field_mappings),
        "equipment_points_count": len(equipment_points),
        "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def batch_export_to_word(
    records: list,
    output_dir: str = "/tmp",
) -> dict:
    """
    批量导出Word文档。
    
    Args:
        records: 记录列表
        output_dir: 输出目录
    
    Returns:
        批量导出结果
    """
    if not HAS_DOCX:
        return {"ok": False, "error": "python-docx 未安装"}
    
    results = []
    success_count = 0
    
    for record in records:
        result = export_record_to_word(
            record_type=record.get("record_type", ""),
            record_data=record.get("record_data", {}),
            device_type=record.get("device_type", ""),
            device_tag=record.get("device_tag", ""),
            workshop=record.get("workshop", ""),
            output_path=os.path.join(output_dir, f"{record.get('record_type', 'doc')}_{datetime.datetime.now().strftime('%Y%m%d%H%M%S')}.docx"),
        )
        results.append(result)
        if result.get("ok"):
            success_count += 1
    
    return {
        "ok": True,
        "total": len(records),
        "success_count": success_count,
        "fail_count": len(records) - success_count,
        "results": results,
    }


def _get_field_mapping(record_type: str) -> list:
    """获取字段映射。"""
    mappings = {
        "construction_log": [
            ("date", "日期"), ("weather", "天气"), ("temperature", "气温"),
            ("workshop", "施工部位"), ("construction_content", "施工内容"),
            ("equipment_used", "使用机械"), ("personnel", "出勤人数"),
            ("progress", "进度情况"), ("quality", "质量情况"),
            ("safety", "安全情况"), ("problems", "存在问题"),
            ("next_plan", "明日计划"), ("recorder", "记录人"),
        ],
        "unboxing_record": [
            ("date", "检验日期"), ("workshop", "施工区域"),
            ("equipment_name", "设备名称"), ("equipment_model", "规格型号"),
            ("equipment_tag", "位号"), ("manufacturer", "制造厂家"),
            ("contract_no", "合同号"), ("packing_condition", "包装情况"),
            ("appearance", "外观检查"), ("parts_check", "零部件清点"),
            ("documents", "随机文件"), ("defects", "缺损情况"),
            ("conclusion", "检验结论"), ("handover_person", "移交人"),
            ("receiver", "接收人"), ("supervisor", "监理"),
        ],
        "concealment_record": [
            ("date", "验收日期"), ("workshop", "施工区域"),
            ("equipment_name", "设备名称"), ("equipment_tag", "位号"),
            ("concealment_part", "隐蔽部位"), ("construction_content", "施工内容"),
            ("material", "材料规格"), ("quality_check", "质量检查"),
            ("test_results", "试验结果"), ("conclusion", "验收结论"),
            ("construction_unit", "施工单位"), ("constructor", "施工负责人"),
            ("quality_inspector", "质检员"), ("supervisor", "监理工程师"),
            ("owner_rep", "建设单位"),
        ],
        "installation_record": [
            ("date", "安装日期"), ("workshop", "施工区域"),
            ("equipment_name", "设备名称"), ("equipment_model", "规格型号"),
            ("equipment_tag", "位号"), ("location", "安装位置"),
            ("elevation", "设备标高"), ("levelness", "水平度"),
            ("verticality", "垂直度"), ("alignment", "同轴度"),
            ("clearance", "关键间隙"), ("bolt_torque", "螺栓力矩"),
            ("grouting", "二次灌浆"), ("quality_standard", "质量标准"),
            ("conclusion", "安装结论"), ("constructor", "施工负责人"),
            ("quality_inspector", "质检员"),
        ],
        "trial_run_record": [
            ("date", "试运转日期"), ("workshop", "施工区域"),
            ("equipment_name", "设备名称"), ("equipment_model", "规格型号"),
            ("equipment_tag", "位号"), ("trial_type", "试运转类型"),
            ("duration", "运转时长"), ("current", "电流(A)"),
            ("voltage", "电压(V)"), ("temperature", "轴承温度(℃)"),
            ("vibration", "振动值(mm/s)"), ("noise", "噪声(dB)"),
            ("pressure", "压力(MPa)"), ("flow", "流量(m³/h)"),
            ("lubrication", "润滑情况"), ("sealing", "密封情况"),
            ("abnormal", "异常情况"), ("conclusion", "试运转结论"),
            ("operator", "操作人"), ("recorder", "记录人"), ("supervisor", "监护人"),
        ],
        "safety_check": [
            ("date", "检查日期"), ("workshop", "检查区域"),
            ("check_type", "检查类型"), ("check_content", "检查内容"),
            ("high_place", "高处作业"), ("lifting", "起重吊装"),
            ("hot_work", "动火作业"), ("confined_space", "受限空间"),
            ("electricity", "临时用电"), ("machinery", "机械设备"),
            ("ppe", "防护用品"), ("hazards", "发现隐患"),
            ("rectification", "整改措施"), ("rectification_person", "整改责任人"),
            ("inspector", "检查人"), ("checked_person", "被检查人"),
        ],
        "lifting_record": [
            ("date", "吊装日期"), ("workshop", "吊装区域"),
            ("equipment_name", "吊装设备"), ("equipment_weight", "设备重量(t)"),
            ("equipment_size", "设备尺寸"), ("lifting_height", "吊装高度(m)"),
            ("crane_type", "吊车类型"), ("crane_working_radius", "作业半径(m)"),
            ("lifting_method", "吊装方法"), ("lifting_points", "吊点设置"),
            ("slings", "索具配置"), ("safety_factor", "安全系数"),
            ("wind_speed", "风速"), ("trial_lift", "试吊情况"),
            ("lifting_process", "吊装过程"), ("positioning", "就位情况"),
            ("abnormal", "异常情况"), ("conclusion", "吊装结论"),
            ("commander", "指挥"), ("crane_operator", "司机"),
            ("signalman", "信号工"), ("safety_officer", "安全员"),
        ],
        "welding_record": [
            ("date", "焊接日期"), ("workshop", "施工区域"),
            ("welding_part", "焊接部位"), ("material", "母材材质"),
            ("welding_process", "焊接方法"), ("welding_material", "焊材牌号"),
            ("welding_material_batch", "焊材质保书"), ("preheat_temp", "预热温度"),
            ("interpass_temp", "层间温度"), ("current", "焊接电流(A)"),
            ("voltage", "焊接电压(V)"), ("welding_speed", "焊接速度"),
            ("welding_layers", "焊接层数"), ("post_weld_heat", "焊后热处理"),
            ("appearance_check", "外观检查"), ("ndt", "无损检测"),
            ("ndt_result", "检测结果"), ("welder", "焊工"),
            ("inspector", "检查员"),
        ],
    }
    return mappings.get(record_type, [])


def _get_equipment_points(record_type: str, device_type: str) -> list:
    """获取设备专用要点。"""
    if not device_type:
        return []
    
    try:
        from . import mining_field_record as _mfr
        points_result = _mfr.get_equipment_record_points(device_type)
        if points_result.get("ok"):
            point_key_map = {
                "unboxing_record": "unboxing",
                "installation_record": "installation",
                "concealment_record": "concealment",
                "trial_run_record": "trial_run",
            }
            point_key = point_key_map.get(record_type)
            if point_key and point_key in points_result["points"]:
                return points_result["points"][point_key]
    except Exception:
        pass
    
    return []
